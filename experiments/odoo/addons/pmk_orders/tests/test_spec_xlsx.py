# -*- coding: utf-8 -*-
"""«Печать → Расчёт (Excel)» (21.10.2026): три листа, для клиента — один."""
import io
import json

import openpyxl

from odoo import Command
from odoo.tests import tagged

from .test_step_z2_invoice import Z2Common

SNAP = json.dumps({"start": {"x": 0, "y": 0}, "segs": [{"len": 120, "dir": 0}, {"len": 80, "dir": 90}]})


@tagged("post_install", "-at_install")
class TestSpecXlsx(Z2Common):

    def _book(self, spec, client=False):
        return openpyxl.load_workbook(io.BytesIO(spec._pmk_xlsx(client=client)))

    def _values(self, ws):
        return [v for row in ws.iter_rows(values_only=True) for v in row if v not in (None, "")]

    def test_three_sheets_with_composition_dobor_vent(self):
        deal = self._deal()
        spec = self._spec(deal, products=[("Опора скользящая", 8, 1400.0)], cat_dobor=True, cat_vent=True)
        product = spec.product_ids
        profile = self.env["pmk.metal.profile"].search([], limit=1)
        sheet = self.env["pmk.metal.sheet"].search([], limit=1)
        lines = []
        if profile:
            lines.append(Command.create({"detail_name": "1", "calc_mode": "linear", "profile_id": profile.id,
                                         "length_mm": 200, "qty": 1}))
        if sheet:
            lines.append(Command.create({"detail_name": "2", "calc_mode": "sheet", "sheet_id": sheet.id,
                                         "a_mm": 150, "b_mm": 600, "qty": 2}))
        product.write({"line_ids": lines})
        coat = self.env.ref("pmk_calc.coating_ral8017")
        spec.dobor_order_ids.sudo().write({"line_ids": [Command.create({
            "title": "Капельник", "coating_id": coat.id, "plank_length": 2500.0, "qty": 20,
            "profile_snapshot_json": SNAP})]})
        spec.vent_order_ids.sudo().write({"line_ids": [Command.create({
            "width": 500, "length": 1000, "qty": 2, "coating_id": coat.id, "price_unit": 1500.0})]})

        book = self._book(spec)
        self.assertEqual(book.sheetnames, ["Изделия", "Состав", "Металл и деньги"])
        products = self._values(book["Изделия"])
        self.assertIn("Опора скользящая", products)
        self.assertIn("Доборка", products)
        self.assertIn("Зонты", products)
        self.assertIn(deal.display_name, products, "сделка — в шапке внутреннего файла")
        composition = self._values(book["Состав"])
        for title in ("Конструкции", "Зонты", "Доборка", "Металл, ₽/т"):
            self.assertIn(title, composition)
        if sheet:
            row = next(r for r in book["Состав"].iter_rows(values_only=True) if r[1] == "2")
            self.assertTrue(row[2].endswith("х150"), "лист — «толщина х ширина», длина отдельно: %s" % row[2])
            self.assertEqual(row[3], 600)
            line = product.line_ids.filtered(lambda l: l.calc_mode == "sheet")
            self.assertEqual(row[10], round(line.price_ton) if line.price_state == "ok" else "нет цены",
                             "₽/т — цена, которую подставила система, а не обратный счёт")
        self.assertIn("Деньги по изделиям", self._values(book["Металл и деньги"]))

    def test_client_file_only_products_without_deal(self):
        deal = self._deal()
        spec = self._spec(deal, products=[("Секция ограждения ОГ-1", 12, 4500.0)])
        book = self._book(spec, client=True)
        self.assertEqual(book.sheetnames, ["Изделия"])
        values = self._values(book["Изделия"])
        self.assertIn("Секция ограждения ОГ-1", values)
        self.assertNotIn(deal.display_name, values)
        self.assertNotIn("Сделка", values)
        self.assertEqual(spec._pmk_xlsx_filename(True), "Расчёт %s для клиента.xlsx" % spec.name)

    def test_print_menu_and_action(self):
        spec = self._spec(self._deal())
        keys = [item["key"] for item in spec.pmk_doc_menu()["print"]]
        self.assertIn("xlsx", keys)
        self.assertIn("xlsx_client", keys)
        action = spec.with_user(self.manager).action_pmk_xlsx_client()
        self.assertEqual(action["url"], "/pmk_orders/spec/%s/xlsx?client=1" % spec.id)
