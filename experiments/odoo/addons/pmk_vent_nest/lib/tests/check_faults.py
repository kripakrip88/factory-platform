# -*- coding: utf-8 -*-
"""Проверки на инвариантах и внедрённых неисправностях (без эталонов) — дополнение к check_golden.py.

    python3 tests/check_faults.py            # быстрые (~1,5 мин)
    python3 tests/check_faults.py --full     # + пик памяти на ВЗ-00002 и сетка генератора целиком
    python3 tests/check_faults.py P05 P21    # выборочно

Каждая проверка ловит конкретную поломку: её нашли мутационным тестированием (ревизия 12.10.2026, P01–P19 —
тесты ревизора) и ревизией второго круга (P20–P27). Неисправности внедряются подменой функций (monkeypatch),
проверка требует, чтобы упала ИМЕННО нужная самопроверка или был нужный отказ.
"""
import sys, os, io, re, math, json, time, zipfile, tempfile, threading, subprocess, contextlib, collections, resource
PKG = os.path.dirname(os.path.dirname(os.path.abspath(__file__))); sys.path.insert(0, PKG)
from vent_nesting import api, nest as N, contours as C, lxds_write as LW, lxds_reader as LR

DATA = os.path.join(PKG, "tests", "data")
Z3 = open(os.path.join(DATA, "ВЗ-00003.zip"), "rb").read()
Q3 = {"750x4000": 1, "1000x1200": 1}
zf = zipfile.ZipFile(io.BytesIO(Z3))
D750 = zf.read("vsh_750x4000_t0.65_RAL8017.dxf")
N750 = "vsh_750x4000_t0.65_RAL8017.dxf"


def quiet(fn, *a, **k):
    with contextlib.redirect_stdout(io.StringIO()):
        return fn(*a, **k)


def raises(fn):
    try:
        quiet(fn)
        return False, "ошибки нет"
    except api.VentNestingError as e:
        return True, "VentNestingError: %s" % str(e)[:90]


def xlsx(rows):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Сборка"
    ws.append(["№", "Зонт", "Цвет", "Металл, мм", "Зонтов, шт", "Деталь", "Заготовка, мм", "Деталей на зонт"])
    for r in rows:
        ws.append(r)
    b = io.BytesIO(); wb.save(b); return b.getvalue()


def mkzip(entries):
    b = io.BytesIO(); z = zipfile.ZipFile(b, "w")
    for n, d in entries:
        z.writestr(n, d)
    z.close(); return b.getvalue()


# ---------------------------------------------------------------- проверки
def P01_arc_midpoints():
    """Каждая дуга контура (bulge) проходит через среднюю точку дуги DXF (pm) — знак bulge независимо от площадей."""
    o = api.order_from_zip(Z3, Q3)
    N.set_spacing(2, 2)
    bad = tot = 0
    for it in quiet(o.items):
        pms = [e["pm"] for e in it["edges"] + it["slot_edges"] if e["t"] == "A"]
        c = C.part_contours(it)
        for poly in [c["outer"]] + c["slots"]:
            for i, (x1, y1, b) in enumerate(poly):
                if not b:
                    continue
                x2, y2, _ = poly[(i + 1) % len(poly)]
                mx, my = (x1 + x2) / 2 + b / 2 * (y2 - y1), (y1 + y2) / 2 - b / 2 * (x2 - x1)
                tot += 1
                if min(math.dist((mx, my), pm) for pm in pms) > 0.05:
                    bad += 1
    return bad == 0 and tot > 0, "дуг %d, не через pm %d" % (tot, bad)


def P02_items_follow_qty():
    """Деталей в раскрое = Σ зонтов × деталей на зонт (на заказе с количеством > 1 — в quick-наборе его нет)."""
    o = api.order_from_zip(Z3, {"750x4000": 2, "1000x1200": 3})
    got = collections.Counter(it["hood"] for it in quiet(o.items))
    want = {"750×4000": 2 * 6, "1000×1200": 3 * 4}
    return dict(got) == want, "получено %s, ждём %s" % (dict(got), want)


def P03_duplicates():
    """Один зонт дважды: тем же именем с количеством по имени; другим написанием толщины/цвета."""
    r1 = raises(lambda: api.order_from_dxf([(N750, D750)] * 2, {N750: 1}))
    r2 = raises(lambda: api.order_from_dxf([(N750, D750), ("vsh_750x4000_t0,65_RAL8017.dxf", D750)],
                                           {N750: 1, "vsh_750x4000_t0,65_RAL8017.dxf": 1}))
    r3 = raises(lambda: api.order_from_dxf([(N750, D750), ("vsh_750x4000_t0.65_RAL 8017.dxf", D750)],
                                           {N750: 1, "vsh_750x4000_t0.65_RAL 8017.dxf": 1}))
    return r1[0] and r2[0] and r3[0], "то же имя: %s | t0,65: %s | 'RAL 8017': %s" % (r1[1][:40], r2[1][:40], r3[1][:40])


def P04_lost_part_fault():
    """Внедрённая неисправность: раскрой потерял деталь → самопроверка «по каждому зонту = заказ» обязана упасть."""
    o = api.order_from_zip(Z3, Q3)
    orig = N.load_order
    N.load_order = lambda paths, qty: orig(paths, qty)[:-1]
    try:
        f = quiet(api.build_lxds, o, "stock", material="оц|0.65")[0]
    finally:
        N.load_order = orig
    failed = [c[0] for c in f["checks"] if not c[1]]
    return (not f["ok"]) and "деталей по каждому зонту = заказ" in failed, "ok=%s, упали %s" % (f["ok"], failed)


def P05_slot_arc_flip_fault():
    """Внедрённая неисправность: сборщик пишет дуги наоборот, сверка размещений отключена →
    проверка площадей контуров обязана упасть."""
    o = api.order_from_zip(Z3, Q3)
    lw0, cp0 = LW._lw, LW.check_placed
    LW._lw = lambda pts, outer=0: lw0([(x, y, -b) for x, y, b in pts], outer)
    LW.check_placed = lambda shapes, cont: True
    try:
        f = quiet(api.build_lxds, o, "stock", material="оц|0.65")[0]
    finally:
        LW._lw, LW.check_placed = lw0, cp0
    failed = [c[0] for c in f["checks"] if not c[1]]
    return (not f["ok"]) and "контуры деталей = раскрой (площадь каждого контура)" in failed, "ok=%s, упали %s" % (f["ok"], failed)


def P06_cutlen_fault():
    """Внедрённая неисправность: длина реза модели на 1% больше → «длина реза = раскрой» обязана упасть."""
    o = api.order_from_zip(Z3, Q3)
    l0 = C.length
    C.length = lambda p: l0(p) * 1.01
    try:
        f = quiet(api.build_lxds, o, "stock", material="оц|0.65")[0]
    finally:
        C.length = l0
    failed = [c[0] for c in f["checks"] if not c[1]]
    return (not f["ok"]) and "длина реза = раскрой" in failed, "ok=%s, упали %s" % (f["ok"], failed)


def P07_pph_mismatch():
    """Ведомость: «Деталей на зонт» 5, генератор даёт 6 → ошибка."""
    x = xlsx([[1, "вш 750х4000", "RAL 8017", 0.65, 1, "Крышка", "", 2], [None, None, None, None, None, "Панель", "", 3]])
    return raises(lambda: api.order_from_zip(mkzip([(N750, D750), ("X_сборка.xlsx", x)])))


def P08_assembly_two_rows():
    """Ведомость: один зонт двумя строками (1 + 2) → 3 зонта."""
    x = xlsx([[1, "вш 750х4000", "RAL 8017", 0.65, 1, "Крышка", "", 2], [None, None, None, None, None, "Панель", "", 4],
              [2, "вш 750х4000", "RAL 8017", 0.65, 2, "Крышка", "", 2], [None, None, None, None, None, "Панель", "", 4]])
    o = api.order_from_zip(mkzip([(N750, D750), ("X_сборка.xlsx", x)]))
    q = o.hoods()[0]["qty"]
    return q == 3, "qty=%s" % q


def P09_qty_over_max():
    return raises(lambda: api.order_from_dxf([(N750, D750)], {N750: api.MAX_HOODS + 1}))


def P10_material_by_colour_name():
    """material — ключ ИЛИ название цвета (README): «оцинковка» и «RAL 8017»."""
    o = api.order_from_zip(Z3, Q3)
    a = [f["key"] for f in quiet(api.build_lxds, o, "stock", material="оцинковка")]
    b = [f["key"] for f in quiet(api.build_lxds, o, "stock", material="RAL 8017")]
    return a == ["оц|0.65"] and b == ["ral8017|0.65"], "оцинковка→%s, RAL 8017→%s" % (a, b)


def P11_sheet_dxf_readback():
    """DXF листов: окружности = отверстия деталей листа (радиус = d/2, число)."""
    o = api.order_from_zip(Z3, Q3)
    p = quiet(api.plan, o, png=False, sheet_dxf=True)
    holes = collections.Counter(round(h[2] / 2, 3) for it in quiet(o.items) for h in it["holes"])
    got = collections.Counter()
    for name, txt in p["dxf"].items():
        if "_склад_" not in name:
            continue
        L = txt.split("\n")
        for i, v in enumerate(L):
            if v == "CIRCLE":
                got[round(float(L[L.index("40", i) + 1]), 3)] += 1
    return got == holes, "окружности %s, отверстия %s" % (dict(got), dict(holes))


def P12_cli_smoke():
    """CLI lxds с --gap/--edge: код 0, имя файла с «_зазор10_край10», в выводе все проверки ✓."""
    with tempfile.TemporaryDirectory() as tmp:
        pr = subprocess.run([sys.executable, "-B", "-m", "vent_nesting.cli", "lxds", os.path.join(DATA, "ВЗ-00003.zip"),
                             "--qty", "750x4000=1", "--qty", "1000x1200=1", "--source", "stock", "--material", "оц|0.65",
                             "--gap", "10", "--edge", "10", "--out", tmp], cwd=PKG, capture_output=True, text=True,
                            env=dict(os.environ, PYTHONPATH=os.environ.get("PYTHONPATH", "")))
        files = os.listdir(tmp)
    ok = pr.returncode == 0 and any("_зазор10_край10" in f for f in files)
    return ok, "rc=%d файлы %s %s" % (pr.returncode, files, pr.stderr[-200:])


def P13_sves150_is_default():
    """sves=150 ≡ sves=None: те же листы, имя файла без «_свес150»."""
    o = api.order_from_zip(Z3, Q3)
    a = quiet(api.plan, o, sves=150, png=False)["materials"]
    b = quiet(api.plan, o, png=False)["materials"]
    fn = quiet(api.build_lxds, o, "stock", sves=150, material="оц|0.65")[0]["filename"]
    return a == b and "свес" not in fn, "plan равны: %s; файл %s" % (a == b, fn)


def P14_gcache_bounded():
    """Кэш спектров ограничен: после plan(ВЗ-00003) учёт байт сходится и ≤ 128 МБ (без потолка — 288 МБ)."""
    o = api.order_from_zip(Z3, Q3)
    quiet(api.plan, o, png=False)
    real = sum(G.nbytes for G in N._gcache.values())
    return real == N._gbytes[0] and real <= 128 * 2 ** 20, "кэш %d шт, %.0f МБ (учёт %.0f МБ), _GCAP=%s" % (
        len(N._gcache), real / 2 ** 20, N._gbytes[0] / 2 ** 20, N._GCAP)


def P14b_peak_memory():
    """Пик памяти plan(ВЗ-00002) в отдельном процессе < 450 МБ (полный набор; предел контейнера 2,5 ГБ, Odoo ~0,9 ГБ)."""
    code = ("import sys,resource;sys.path.insert(0,%r);from vent_nesting import api;"
            "o=api.order_from_zip(open(%r,'rb').read());api.plan(o,png=False);"
            "r=resource.getrusage(resource.RUSAGE_SELF).ru_maxrss;print(r/2**20 if sys.platform=='darwin' else r/1024)"
            % (PKG, os.path.join(DATA, "ВЗ-00002.zip")))
    pr = subprocess.run([sys.executable, "-B", "-c", code], capture_output=True, text=True)
    mb = float(pr.stdout.strip().splitlines()[-1])
    return mb < 450, "пик %.0f МБ" % mb


def P15_build_honours_sves_quick():
    """build_lxds со свесом 140 режет меньше, чем со 150 (ВЗ-00003, склад, оцинковка)."""
    o = api.order_from_zip(Z3, Q3)
    def cut(sv):                         # точная длина реза по прочитанной программе (в отчёте — округлённая до 0,1 м)
        f = quiet(api.build_lxds, o, "stock", sves=sv, material="оц|0.65")[0]
        r, ctx = quiet(LR.build, f["bytes"], name=f["filename"], thickness=0.65, min_side=None)
        return sum(d["sheets"] * sum(LR.elen(e) for e in d["_ents"]) for d in ctx["lays"]) / 1000
    a, b = cut(None), cut(140)
    return b < a - 0.01, "рез 150: %.3f м, 140: %.3f м" % (a, b)


def P16_dxf_order_qty_gt1():
    """order_from_dxf с количеством > 1: hoods() и items() его сохраняют."""
    o = api.order_from_dxf([(N750, D750)], {N750: 7})
    n = len(quiet(o.items))
    return o.hoods()[0]["qty"] == 7 and n == 42, "qty=%s, деталей %d" % (o.hoods()[0]["qty"], n)


def P17_rows_after_itogo():
    """Ведомость: строки ПОСЛЕ «ИТОГО» (примечания с числом в колонке H) не прибавляются к последнему зонту."""
    x = xlsx([[1, "вш 750х4000", "RAL 8017", 0.65, 1, "Крышка", "", 2], [None, None, None, None, None, "Панель", "", 4],
              ["ИТОГО", None, None, None, 1, None, None, None], ["Примечание: запас", None, None, None, None, None, None, 5]])
    try:
        o = api.order_from_zip(mkzip([(N750, D750), ("X_сборка.xlsx", x)]))
    except api.VentNestingError as e:
        return False, "VentNestingError: %s" % str(e)[:100]
    return o.hoods()[0]["parts_per_hood"] == 6, "parts_per_hood=%s" % o.hoods()[0]["parts_per_hood"]


def P17b_itogo_row_with_parts():
    """Ведомость: число в колонке H самой строки «ИТОГО» не прибавляется к последнему зонту."""
    x = xlsx([[1, "вш 750х4000", "RAL 8017", 0.65, 1, "Крышка", "", 2], [None, None, None, None, None, "Панель", "", 4],
              ["ИТОГО", None, None, None, 1, None, None, 6]])
    try:
        o = api.order_from_zip(mkzip([(N750, D750), ("X_сборка.xlsx", x)]))
    except api.VentNestingError as e:
        return False, "VentNestingError: %s" % str(e)[:100]
    return o.hoods()[0]["parts_per_hood"] == 6, "parts_per_hood=%s" % o.hoods()[0]["parts_per_hood"]


def P18_qty_twice():
    """Количество одному DXF двумя ключами (имя и размер) → ошибка, а не «последний выиграл»."""
    return raises(lambda: api.order_from_dxf([(N750, D750)], {N750: 1, "750x4000": 2}))


def P19_two_thicknesses():
    """Один цвет в двух толщинах (0,65 ×1 и 0,5 ×2): две программы, обе ok, каждая со своим числом деталей."""
    n5 = "vsh_750x4000_t0.5_RAL8017.dxf"
    o = api.order_from_dxf([(N750, D750), (n5, D750)], {N750: 1, n5: 2})
    fs = quiet(api.build_lxds, o, "stock")
    det = [(f["key"], f["ok"], [c[2] for c in f["checks"] if c[0].startswith("деталей по каждому")]) for f in fs]
    return len(fs) == 2 and all(f["ok"] for f in fs), str(det)[:220]


# ---------------------------------------------------------------- второй круг ревизии
from vent_nesting import hood_sves as V


def _gen(W, L, q=1, t="0.65", col="RAL8017"):
    n = "vsh_%dx%d_t%s_%s.dxf" % (W, L, t, col)
    return n, V.to_dxf(V.hood(W, L)).encode()


def P20_generator_grid_contours():
    """Все детали генератора сцепляются в контуры (сетка Ш×Д×свес). Раньше клетка округления рвала пазы у Ш=1310/1350."""
    from vent_nesting import contours as CC
    full = "--full" in sys.argv
    Ws = range(200, 1351, 10 if full else 50)
    Ls = list(range(300, 6001, 250)) + [2300, 3050, 5800, 6000]
    bad = n = 0
    with tempfile.TemporaryDirectory() as d:
        for W in Ws:
            for L in Ls:
                if W > L:
                    continue
                for sv in ((150, 145, 135, 125) if full else (150, 125)):
                    try:
                        g = V.hood(W, L, sves=sv)
                    except ValueError:
                        continue
                    p = os.path.join(d, "x.dxf"); open(p, "w").write(V.to_dxf(g))
                    for it in N.load_order([(p, "vsh_%dx%d_t0.65_RAL8017.dxf" % (W, L))], {}):
                        n += 1
                        try:
                            CC.part_contours(it)
                        except ValueError:
                            bad += 1
    return bad == 0 and n > 0, "деталей %d, не сцепилось %d" % (n, bad)


def P21_previously_crashing_sizes():
    """Размеры, на которых сборка падала (сцепка контуров): 1350x3000, 1310x2000, 750x2300, 900x2300 — ok=True."""
    res = []
    for W, L in ((1350, 3000), (1310, 2000), (750, 2300), (900, 2300)):
        n, d = _gen(W, L)
        f = quiet(api.build_lxds, api.order_from_dxf([(n, d)], {n: 1}), "stock")[0]
        res.append((W, L, f["ok"]))
    return all(r[2] for r in res), str(res)


def P22_sves_changes_parts():
    """Свес 145 у 500x4650 даёт 4 детали на зонт вместо 5 — программа ok=True (было ложное «НЕ ВЫДАВАТЬ»)."""
    n, d = _gen(500, 4650)
    o = api.order_from_dxf([(n, d)], {n: 2})
    f = quiet(api.build_lxds, o, "stock", sves=145)[0]
    det = [c[2] for c in f["checks"] if c[0].startswith("деталей по каждому")]
    return f["ok"] and o.expected_per_hood("RAL 8017", 0.65, 145) == {"500×4650": 8}, "ok=%s %s" % (f["ok"], det)


def P23_threads_order():
    """Заказы создаются из нескольких потоков: правильный DXF всегда принят, DXF с лишними отверстиями вне деталей —
    всегда отвергнут (раньше перехват stdout путал потоки)."""
    zf2 = zipfile.ZipFile(io.BytesIO(open(os.path.join(DATA, "ВЗ-00002.zip"), "rb").read()))
    good = [(n, zf2.read(n)) for n in zf2.namelist() if n.endswith(".dxf")]
    gq = {n: 1 for n, _ in good}
    txt = D750.decode()
    extra = "".join("0\nCIRCLE\n8\n0\n10\n%d\n20\n0\n30\n0\n40\n3.5\n" % (90000 + 20 * i) for i in range(20))
    badd = txt.replace("0\nENDSEC\n0\nEOF", extra + "0\nENDSEC\n0\nEOF").encode()
    stats = collections.Counter()

    def w_good():
        for _ in range(6):
            try:
                api.order_from_dxf(good, gq); stats["good_ok"] += 1
            except api.VentNestingError:
                stats["good_REJECTED"] += 1

    def w_bad():
        for _ in range(6):
            try:
                api.order_from_dxf([(N750, badd)], {N750: 1}); stats["bad_ACCEPTED"] += 1
            except api.VentNestingError:
                stats["bad_rejected"] += 1
    th = [threading.Thread(target=w_good) for _ in range(3)] + [threading.Thread(target=w_bad) for _ in range(2)]
    [t.start() for t in th]; [t.join() for t in th]
    real_stdout = sys.stdout is sys.__stdout__
    return stats["good_REJECTED"] == 0 and stats["bad_ACCEPTED"] == 0 and real_stdout, "%s, stdout цел: %s" % (dict(stats), real_stdout)


def P24_assembly_layout_v161():
    """Ведомость текущего ERP (v1.61+): «Зонт, мм» = габарит, «Шахта, мм» = «4000 × 750» (длина × ширина)."""
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Сборка"
    ws.append(["Сборочная ведомость вентзонтов — заказ X"]); ws.append([]); ws.append([])
    ws.append(["№", "Зонт, мм", "Шахта, мм", "Цвет", "Металл, мм", "Зонтов, шт", "Деталь", "Заготовка, мм", "Деталей на зонт", "Деталей всего"])
    ws.append([1, "4300 × 1050", "4000 × 750", "RAL 8017 Шоколад", 0.65, 3, "Крышка торцевая", "", 2, 6])
    ws.append([None, None, None, None, None, None, "Панель торцевая", "", 2, 6])
    ws.append([None, None, None, None, None, None, "Панель средняя", "", 2, 6])
    ws.append(["ИТОГО", None, None, None, None, 3, None, None, None, None])
    b = io.BytesIO(); wb.save(b)
    o = api.order_from_zip(mkzip([(N750, D750), ("X_сборка.xlsx", b.getvalue())]))
    h = o.hoods()[0]
    return h["qty"] == 3 and h["parts_per_hood"] == 6 and h["size"] == "750x4000", str(h)


def P24b_assembly_v161_merged():
    """Ведомость как пишет ERP v1.61 (xlsxwriter): ячейки зонта объединены на высоту его деталей, ИТОГО объединено
    до «Металл, мм», «Деталей всего» — формулой; два зонта, один размер в двух цветах."""
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Сборка"
    ws.append(["Сборочная ведомость вентзонтов — заказ X"]); ws.append(["Клиент: —"]); ws.append(["Изделие / объект: —"])
    ws.append(["Дата: 12.10.2026"]); ws.append([])
    ws.append(["№", "Зонт, мм", "Шахта, мм", "Цвет", "Металл, мм", "Зонтов, шт", "Деталь", "Заготовка, мм", "Деталей на зонт",
               "Деталей всего", "Лапок на зонт", "Лапок всего", "Клёпок на зонт", "Клёпок всего"])
    r0 = ws.max_row + 1
    for k, (col, q) in enumerate((("RAL 8017 Шоколад", 2), ("", 5))):            # пустой цвет = оцинковка
        top = ws.max_row + 1
        for i, (part, nn) in enumerate((("Крышка торцевая", 2), ("Панель торцевая", 2), ("Панель средняя", 2))):
            row = [k + 1, "4300 × 1050", "4000 × 750", col, 0.65, q] if i == 0 else [None] * 6
            ws.append(row + [part, "", nn, "=I%d*F%d" % (top + i, top), 30, None, 40, None])
        for c in range(1, 7):
            ws.merge_cells(start_row=top, start_column=c, end_row=top + 2, end_column=c)
    ws.append(["ИТОГО", None, None, None, None, 7]); ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=5)
    ws.append([]); ws.append(["Не вошёл: вш 9999 — размер не задан"])
    b = io.BytesIO(); wb.save(b)
    n_oc = "vsh_750x4000_t0.65_оц.dxf"
    o = api.order_from_zip(mkzip([(N750, D750), (n_oc, D750), ("X_сборка.xlsx", b.getvalue())]))
    got = sorted((h["dxf"], h["qty"], h["parts_per_hood"]) for h in o.hoods())
    want = sorted([(N750, 2, 6), (n_oc, 5, 6)])
    return got == want, str(got)


def P25_long_part_plan():
    """Тело 2500 мм (500x2200) не лезет на склад 2500 — план не падает: склад с ошибкой, бухта считана, совет «бухта»."""
    n, d = _gen(500, 2200)
    m = quiet(api.plan, api.order_from_dxf([(n, d)], {n: 1}), png=True)["materials"][0]
    return "error" in m["stock"] and m["coil"].get("kg") and m["recommend"] == "coil", "склад: %s | бухта %s кг" % (
        str(m["stock"].get("error"))[:50], m["coil"].get("kg"))


def P26_bombs():
    """Подделанный размер в zip (deflate 60 МБ под заголовком 10 КБ) и ведомость на 6000 строк — понятный отказ,
    без съедания памяти."""
    import zlib
    raw = b"0\n" * (30 * 2 ** 20)
    b = io.BytesIO(); z = zipfile.ZipFile(b, "w", zipfile.ZIP_DEFLATED)
    z.writestr(N750, raw); z.close()
    data = bytearray(b.getvalue())
    # подделать file_size в локальном заголовке и центральном каталоге
    import struct
    i = data.find(b"PK\x03\x04"); struct.pack_into("<I", data, i + 22, 10000)
    j = data.find(b"PK\x01\x02"); struct.pack_into("<I", data, j + 24, 10000)
    r1 = raises(lambda: api.order_from_zip(bytes(data), {N750: 1}))
    x = xlsx([[1, "вш 750х4000", "RAL 8017", 0.65, 1, "Крышка", "", 2]] + [[None] * 8] * 6000)
    r2 = raises(lambda: api.order_from_zip(mkzip([(N750, D750), ("X_сборка.xlsx", x)])))
    mb = resource.getrusage(resource.RUSAGE_SELF).ru_maxrss / (2 ** 20 if sys.platform == "darwin" else 1024)
    return r1[0] and r2[0], "zip: %s | xlsx: %s | пик процесса %.0f МБ" % (r1[1][:60], r2[1][:60], mb)


def P27_chain_arc_flip_fault():
    """Внедрённая неисправность в САМОЙ сцепке контуров (знак дуги в contours.chain) → самопроверка
    «контуры деталей = исходный DXF» обязана упасть (раньше это проходило: модель сверялась сама с собой)."""
    from vent_nesting import contours as CC
    o = api.order_from_zip(Z3, Q3)
    sw0 = CC._sweep
    CC._sweep = lambda e, fwd: -sw0(e, fwd)
    try:
        f = quiet(api.build_lxds, o, "stock", material="оц|0.65")[0]
    finally:
        CC._sweep = sw0
    failed = [c[0] for c in f["checks"] if not c[1]]
    return (not f["ok"]) and "контуры деталей = исходный DXF (дуги, отверстия)" in failed, "ok=%s, упали %s" % (f["ok"], failed)


def P28_wrong_types():
    """Неверные типы от интегратора — понятная ошибка, не сырое исключение."""
    cases = [lambda: api.order_from_dxf(None, {}), lambda: api.order_from_dxf([(N750,)], {}),
             lambda: api.plan(None), lambda: api.build_lxds(None, "stock"),
             lambda: api.sves_table(api.order_from_dxf([(N750, D750)], {N750: 1}), sves_list=[]),
             lambda: api.sves_table(api.order_from_dxf([(N750, D750)], {N750: 1}), sves_list=None),
             lambda: api.build_lxds(api.order_from_zip(Z3, Q3), "stock", material="ral8017|0.7")]
    res = [raises(c)[0] for c in cases]
    return all(res), str(res)


def P29_timezone_independent():
    """Программа, собранная в процессе с поясом UTC (контейнер на сервере), проходит все самопроверки:
    время сохранения — заводское UTC+10, не пояс машины (на сервере все программы были ok=False)."""
    code = ("import sys;sys.path.insert(0,%r);from vent_nesting import api;"
            "o=api.order_from_zip(open(%r,'rb').read(),{'750x4000':1,'1000x1200':1});"
            "f=api.build_lxds(o,'stock',material='оц|0.65')[0];print(f['ok'],[c[0] for c in f['checks'] if not c[1]])"
            % (PKG, os.path.join(DATA, "ВЗ-00003.zip")))
    outs = []
    for tz in ("UTC", "Asia/Vladivostok", "Europe/Moscow"):
        pr = subprocess.run([sys.executable, "-B", "-c", code], capture_output=True, text=True, env=dict(os.environ, TZ=tz))
        outs.append((tz, (pr.stdout.strip() or pr.stderr.strip()[-120:])))
    return all(o.startswith("True") for _, o in outs), str(outs)


def P30_other_machine_ulp():
    """«Другая машина»: координаты деталей, по которым идёт раскрой, сдвинуты на несколько единиц последнего знака
    (как дают другой libm/процессор) — раскрой обязан не измениться. 12.10 сервер разложил ВЗ-00003 оцинковку иначе
    из-за последнего знака площади двух одинаковых крышек."""
    import random
    o = api.order_from_zip(Z3, Q3)
    base = quiet(api.plan, o, png=False)["materials"]
    orig = N.load_order
    out = []
    for seed in (1, 2, 3):
        rnd = random.Random(seed)
        def noisy(paths, qty, _o=orig, _r=rnd):
            items = _o(paths, qty)
            for it in items:
                it["poly"] = [(x * (1 + _r.choice((-3, -1, 1, 3)) * 2.0 ** -52), y * (1 + _r.choice((-3, -1, 1, 3)) * 2.0 ** -52))
                              for x, y in it["poly"]]
            return items
        N.load_order = noisy
        try:
            m = quiet(api.plan, o, png=False)["materials"]
        finally:
            N.load_order = orig
        same = all(a[k]["placements"] == b[k]["placements"] and a[k]["lengths_mm"] == b[k]["lengths_mm"]
                   for a, b in zip(base, m) for k in ("stock", "coil"))
        out.append(same)
    return all(out), "совпало при шуме: %s" % out


CHECKS = [v for k, v in sorted(globals().items()) if re.match(r"P\d\d[a-z]?_", k)]
SLOW = {"P14b_peak_memory"}

if __name__ == "__main__":
    names = [a for a in sys.argv[1:] if not a.startswith("--")]
    full = "--full" in sys.argv
    bad = n = 0
    for fn in CHECKS:
        if names and fn.__name__ not in names and fn.__name__.split("_")[0] not in names:
            continue
        if fn.__name__ in SLOW and not full and not names:
            continue
        n += 1; t = time.time()
        try:
            ok, det = fn()
        except Exception as e:
            ok, det = False, "ИСКЛЮЧЕНИЕ %s: %s" % (type(e).__name__, str(e)[:150])
        bad += not ok
        print("%s %-30s %4.0f с  %s" % ("✓" if ok else "✗", fn.__name__, time.time() - t, det), flush=True)
    print("\nпроверок %d, с ошибками %d" % (n, bad))
    sys.exit(1 if bad else 0)

