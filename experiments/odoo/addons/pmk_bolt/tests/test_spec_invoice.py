# -*- coding: utf-8 -*-
"""Болты в расчёте и счёте (21.10.2026): категория, вкладка, сводка, счёт из расчёта."""
from unittest.mock import patch

from odoo.tests import tagged

from odoo.addons.pmk_orders.tests.test_step_z2_invoice import Z2Common


@tagged("post_install", "-at_install")
class TestBoltInSpecAndInvoice(Z2Common):

    def setUp(self):
        super().setUp()
        patcher = patch.object(type(self.env["pmk.bolt.line"]), "_pmk_ref_price_ton",
                               lambda line, record, grade=None: 90000.0)
        patcher.start()
        self.addCleanup(patcher.stop)

    def _bolt_spec(self, deal=None):
        deal = deal or self._deal()
        spec = self._spec(deal, products=[("Ферма Ф-1", 2, 1000.0)])
        spec.cat_bolt = True
        rows = spec.bolt_rows()["rows"]
        spec.bolt_save_row(rows[0]["id"], {"name": "бф 2.1 м30 600 09г2с 16шт гайка 45"})
        line = spec.bolt_calc_ids.line_ids.filtered("name")
        line.part_ids.filtered(lambda p: p.auto_key == "washer")._pmk_user_write({"price_unit": 30})
        return deal, spec, line

    def test_category_tab_and_summary(self):
        _deal, spec, line = self._bolt_spec()
        self.assertEqual(len(spec.bolt_calc_ids), 1)
        self.assertEqual(spec.bolt_calc_ids.partner_id, spec.partner_id)
        self.assertEqual(line.qty, 16)
        bolt = next(r for r in spec.pmk_summary["rows"] if r["source"] == "bolt")
        self.assertIn("тип 2.1 М-30х600", bolt["name"])
        self.assertEqual([d["name"][:5] for d in bolt["details"]], ["Круг ", "Гайка", "Шайба", "Плита"])
        # Выключили — черновик: в сводку не идёт, документ цел.
        spec.cat_bolt = False
        self.assertTrue(spec.bolt_calc_ids)
        self.assertFalse([r for r in spec.pmk_summary["rows"] if r["source"] == "bolt"])
        # Копия расчёта — с копией болтов.
        spec.cat_bolt = True
        copy = spec.copy()
        self.assertEqual(copy.bolt_calc_ids.line_ids.filtered("name").mapped("name"), [line.name])

    def test_invoice_from_spec_follows_bolt(self):
        deal, spec, line = self._bolt_spec()
        spec.with_user(self.manager).action_pmk_invoice()
        invoice = self._invoices(deal)
        bolt_line = invoice.order_line.filtered("pmk_bolt_line_id")
        self.assertEqual(bolt_line.name, line.name)
        self.assertEqual((bolt_line.product_uom_qty, bolt_line.price_unit), (16, line.price_unit))
        # Цена болта поменялась — черновик следом.
        line.price_unit = 2500
        self.assertEqual(bolt_line.price_unit, 2500)
        line.qty = 20
        self.assertEqual(bolt_line.product_uom_qty, 20)
        # Без цены — в «без цены клиенту», не в счёте.
        line.price_unit = 0
        self.assertFalse(invoice.order_line.filtered("pmk_bolt_line_id"))
        self.assertIn(line.short_name, invoice.pmk_skip_text or "")

    def test_invoice_menu_bolts(self):
        deal = self._deal()
        spec = self._spec(deal, products=[("Секция ограждения ОГ-1", 2, 4500.0)])
        spec.with_user(self.manager).action_pmk_invoice()
        invoice = self._invoices(deal)
        keys = [i["key"] for i in invoice.pmk_doc_menu()["create"]]
        # Задача 318 (25.10.2026): болты — раздел расчёта, в «Создать ▾» счёта их нет.
        self.assertNotIn("bolts", keys)
        action = invoice.action_pmk_doc_bolts()
        self.assertTrue(spec.cat_bolt)
        self.assertEqual(action.get("res_id"), spec.id)

    def test_header_with_bolts(self):
        """Шапка расчёта (22.10.2026): цена, себестоимость, вес и «нет цены» — с болтами."""
        _deal, spec, line = self._bolt_spec()
        products = sum(spec.product_ids.mapped("price_customer_total"))
        self.assertAlmostEqual(spec.price_customer_total, products + line.price_total, places=2)
        self.assertAlmostEqual(spec.cost_all_total, spec.total_cost_fact + line.cost_one * line.qty, places=2)
        self.assertAlmostEqual(spec.margin_amount, spec.price_customer_total - spec.cost_all_total, places=2)
        self.assertEqual(spec.price_uncosted, 0)
        self.assertAlmostEqual(spec.total_weight, sum(spec.product_ids.mapped("weight_total"))
                               + line.mass_one * line.qty, places=2)
        self.assertEqual(spec.total_products, len(spec.product_ids) + 1)
        before = spec.no_price_count
        # Шайба без цены — себестоимость занижена: «нет цены» в шапке с названием.
        line.part_ids.filtered(lambda p: p.auto_key == "washer")._pmk_user_write({"price_unit": 0})
        self.assertEqual(spec.no_price_count, before + 1)
        self.assertTrue(spec.price_incomplete)
        self.assertIn("Шайба", spec.price_missing_text)
        # Болты выключили — шапка без них.
        spec.cat_bolt = False
        self.assertAlmostEqual(spec.price_customer_total, products, places=2)
        self.assertEqual(spec.no_price_count, before)

    def test_xlsx_composition_bolts(self):
        """«Расчёт (Excel)»: на листе «Состав» — раздел болтов с составом."""
        import io
        from openpyxl import load_workbook
        _deal, spec, line = self._bolt_spec()
        book = load_workbook(io.BytesIO(spec._pmk_xlsx()))
        cells = [str(c.value) for row in book["Состав"].iter_rows() for c in row if c.value]
        self.assertIn("Фундаментные болты", cells)
        self.assertIn(line.name, cells)
        self.assertTrue(any(c.startswith("Круг 30") for c in cells))
        money = [str(c.value) for row in book["Металл и деньги"].iter_rows() for c in row if c.value]
        self.assertTrue(any(c.startswith("Гайка М30") for c in money))

    def test_summary_entry_bolt(self):
        """Ввод в сводке: «бф …» — изделием «Конструкций» (Антон 25.10.2026, вариант 1): шпилька —
        круг, гайки и шайба — метизы; категория болтов не включается."""
        deal = self._deal()
        spec = self._spec(deal, products=[("Ферма Ф-1", 2, 1000.0)])
        res = spec.pmk_summary_add(["бф 1.1 м20 1000 ст3 8шт", "бф м20 8шт"])
        self.assertEqual(res["added"], {"Болты": 1})
        self.assertEqual(len(res["errors"]), 1)
        self.assertFalse(spec.cat_bolt)
        product = spec.product_ids.filtered(lambda p: p.pmk_entry_text == "бф 1.1 м20 1000 ст3 8шт")
        self.assertEqual(product.qty, 8)
        modes = product.line_ids.sorted("sequence").mapped("calc_mode")
        self.assertEqual(modes[0], "linear")
        self.assertIn("fastener", modes)
        stud = product.line_ids.filtered(lambda l: l.calc_mode == "linear")
        self.assertEqual(stud.profile_id.profile_type, "Круг")
        row = next(r for r in spec.pmk_summary["rows"] if r["key"] == "metal-%s" % product.id)
        self.assertEqual(row["source"], "metal")

    def test_bolt_to_spec_konstr(self):
        """«В расчёт» в калькуляторе: болт — изделие «Конструкций» на номенклатуре; цена гаек и
        шайбы — вручную за штуку (в себестоимости); повторный перенос — то же изделие."""
        deal = self._deal()
        spec = self._spec(deal, products=[("Ферма Ф-1", 2, 1000.0)])
        calc = self.env["pmk.bolt.calc"].create({"opportunity_id": deal.id, "partner_id": deal.partner_id.id})
        rows = calc.bolt_rows()["rows"]
        calc.bolt_save_row(rows[0]["id"], {"name": "бф 2.1 м30 600 09г2с 16шт гайка 45"})
        line = calc.line_ids.filtered("name")
        line.part_ids.filtered(lambda p: p.auto_key == "washer")._pmk_user_write({"price_unit": 30})
        calc.spec_id = spec
        action = calc.action_bolt_to_spec()
        self.assertEqual(action["res_id"], spec.id)
        product = line.konstr_product_id
        self.assertEqual((product.spec_id, product.qty), (spec, 16))
        self.assertAlmostEqual(product.price_customer_unit, line.price_unit, places=2)
        stud = next(l for l in product.line_ids if l.calc_mode == "linear")
        self.assertEqual(stud.profile_id.size_label, "d30")
        self.assertEqual(stud.grade_id.name, "09Г2С")
        plate = next(l for l in product.line_ids if l.calc_mode == "sheet")
        self.assertTrue(plate.a_mm > 0 and plate.a_mm == plate.b_mm and plate.sheet_id.thickness_mm > 0)
        nut = next(l for l in product.line_ids if l.calc_mode == "fastener" and l.fastener_id.fastener_type == "nut")
        nuts = int(line.part_ids.filtered(lambda p: p.auto_key == "nut").qty)
        self.assertEqual((nut.fastener_id.size_label, nut.qty, nut.price_state), ("М30", nuts, "manual"))
        self.assertAlmostEqual(nut.price_unit, 45)
        self.assertAlmostEqual(nut.cost_clean_total, 45 * nuts, places=2)
        washer = next(l for l in product.line_ids if l.fastener_id.fastener_type == "washer")
        self.assertAlmostEqual(washer.price_unit, 30)
        # Повторный перенос — то же изделие, без второго.
        calc.action_bolt_to_spec()
        self.assertEqual(line.konstr_product_id, product)
        self.assertEqual(len(spec.product_ids.filtered(lambda p: p.name == line.name)), 1)

    def test_migration_category_to_konstr(self):
        """Миграция 19.0.1.3.0: болты категории «Фундаментные болты» — изделиями «Конструкций»."""
        import importlib.util
        import os
        from odoo.modules.module import get_module_path
        _deal, spec, line = self._bolt_spec()
        path = os.path.join(get_module_path("pmk_bolt"), "migrations", "19.0.1.3.0", "post-migrate.py")
        spec_mod = importlib.util.spec_from_file_location("pmk_bolt_mig_1930", path)
        mod = importlib.util.module_from_spec(spec_mod)
        spec_mod.loader.exec_module(mod)
        mod.migrate(self.env.cr, "19.0.1.2.0")
        spec.invalidate_recordset()
        self.assertFalse(spec.cat_bolt)
        self.assertTrue(line.konstr_product_id)
        self.assertEqual(line.konstr_product_id.spec_id, spec)
        self.assertFalse(any(r["source"] == "bolt" for r in spec.pmk_summary["rows"]))

    def test_invoice_follows_summary_order(self):
        """№312: порядок сводки — порядок строк счёта (болт перетащили наверх)."""
        deal, spec, line = self._bolt_spec()
        spec.with_user(self.manager).action_pmk_invoice()
        invoice = self._invoices(deal)
        names = invoice.order_line.sorted("sequence").mapped("name")
        self.assertEqual(names[-1], line.name, "Не перетаскивали — болты после конструкций.")
        keys = [r["key"] for r in spec.pmk_summary["rows"]]
        bolt_key = "bolt-%s" % line.id
        spec.pmk_summary_order([bolt_key] + [k for k in keys if k != bolt_key])
        names = invoice.order_line.sorted("sequence").mapped("name")
        self.assertEqual(names[0], line.name, "Болт первым и в счёте.")

    def test_grid_move_rows(self):
        """№312: болт — среди болтов; строка состава по ГОСТ в чужой болт не переносится."""
        from odoo.exceptions import UserError
        _deal, spec, line = self._bolt_spec()
        doc = spec.bolt_calc_ids
        second = doc._grid_new_bolt(line, "бф 1.1 м20 1000 ст3 4шт")
        doc.bolt_move_row("b%s" % second.id, "b%s" % line.id, False)
        self.assertEqual(doc._grid_lines().filtered("name")[:1], second)
        nut = line.part_ids.filtered(lambda p: p.auto_key == "nut")
        with self.assertRaises(UserError):
            doc.bolt_move_row("p%s" % nut.id, "b%s" % second.id, False)
        washer = line.part_ids.filtered(lambda p: p.auto_key == "washer")
        doc.bolt_move_row("p%s" % washer.id, "p%s" % nut.id, False)
        self.assertEqual(line.part_ids.sorted(lambda p: (p.sequence, p.id)).mapped("auto_key")[:2], ["stud", "washer"])
